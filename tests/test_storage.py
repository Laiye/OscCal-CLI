"""实际文件持久化回归：唯一文件名、原子更新及元数据追溯。"""

from datetime import datetime
from pathlib import Path

import pytest

from osccal.core import storage


@pytest.fixture
def data_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "DATA_DIR", str(tmp_path))
    return tmp_path


def test_same_timestamp_does_not_overwrite_and_metadata_roundtrips(data_dir, monkeypatch):
    class FrozenDatetime:
        @staticmethod
        def now():
            return datetime(2026, 9, 14, 12, 0, 0)

    monkeypatch.setattr(storage, "datetime", FrozenDatetime)
    metadata = {
        "scpi_errors": 3,
        "status": "completed_with_errors",
        "failures": {"amp": {"type": "ScpiError", "message": "读取失败"}},
        "configurations": {"profile": storage.configuration_snapshot({"name": "测试配置"})},
        "extra": {"保留扩展字段": True},
    }
    first = storage.save_calibration_data({"amp": []}, metadata)
    second = storage.save_calibration_data({"amp": [{"measured": 1}]}, {})
    assert first != second
    assert len(storage.list_calibration_files()) == 2
    loaded = storage.load_calibration_data(first)
    assert loaded["results"] == {"amp": []}
    for key, value in metadata.items():
        assert loaded["metadata"][key] == value
    assert storage.load_calibration_data(second)["metadata"]["scpi_errors"] == 0


@pytest.mark.parametrize("failure_at", ["serialize", "flush", "replace", "interrupt"])
def test_failed_update_keeps_previous_json(data_dir, monkeypatch, failure_at):
    path = storage.save_calibration_data({"amp": [{"measured": 1}]}, {"status": "in_progress"})
    previous = Path(path).read_bytes()

    def fail(*args, **kwargs):
        if failure_at == "interrupt":
            raise KeyboardInterrupt
        raise OSError("模拟磁盘写入失败")

    results = {"amp": [{"measured": object() if failure_at == "serialize" else 2}]}
    if failure_at == "flush":
        monkeypatch.setattr(storage.os, "fsync", fail)
    elif failure_at in {"replace", "interrupt"}:
        monkeypatch.setattr(storage.os, "replace", fail)
    with pytest.raises((TypeError, OSError, KeyboardInterrupt)):
        storage.save_calibration_data(results, {"status": "completed"}, filepath=path)
    assert Path(path).read_bytes() == previous
    assert len(list(data_dir.iterdir())) == 1, "失败写入的临时文件应被清理"


def test_checkpoint_replaces_only_its_own_file(data_dir):
    metadata = {"timestamp": "2026-09-14T12:00:00", "status": "in_progress"}
    path = storage.save_calibration_data({}, metadata)
    other = storage.save_calibration_data({"amp": []}, {})
    other_contents = Path(other).read_bytes()
    metadata.update(status="completed", scpi_errors=4)
    updated = storage.save_calibration_data({"amp": [{"measured": 2}]}, metadata, filepath=path)
    assert updated == path
    assert Path(other).read_bytes() == other_contents
    loaded = storage.load_calibration_data(path)
    assert loaded["metadata"]["timestamp"] == "2026-09-14T12:00:00"
    assert loaded["metadata"]["status"] == "completed"
    assert loaded["metadata"]["scpi_errors"] == 4
    assert loaded["results"]["amp"][0]["measured"] == 2


def test_latest_ignores_temporary_files_and_supports_old_names(data_dir):
    old = data_dir / "calibration_20200101_120000.json"
    old.write_text('{"metadata": {}, "results": {}}', encoding="utf-8")
    assert storage.get_latest_calibration_file() == str(old)
    new = storage.save_calibration_data({}, {})
    (data_dir / ".calibration_unfinished.tmp").write_text('{"metadata":', encoding="utf-8")
    assert storage.get_latest_calibration_file() == new
    assert len(storage.list_calibration_files()) == 2


def test_configuration_hash_tracks_content_and_snapshot_is_independent():
    config = {"name": "配置", "points": [1, 2]}
    snapshot = storage.configuration_snapshot(config)
    assert snapshot == storage.configuration_snapshot({"points": [1, 2], "name": "配置"})
    config["points"].append(3)
    assert snapshot["data"]["points"] == [1, 2]
    assert snapshot["sha256"] != storage.configuration_snapshot(config)["sha256"]


def test_partial_record_status_and_failures_visible_in_show_and_excel(data_dir):
    from click.testing import CliRunner
    from openpyxl import load_workbook

    from osccal.cli import cli
    from osccal.core.export import export_to_excel

    path = storage.save_calibration_data(
        {"amp": []},
        {
            "status": "interrupted",
            "scpi_errors": 2,
            "failures": {"amp": {"type": "ScpiError", "message": "读取失败"}},
        },
    )
    shown = CliRunner().invoke(cli, ["show", "--file", path])
    assert shown.exit_code == 0
    assert "已中断" in shown.output
    assert "读取失败" in shown.output
    output = data_dir / "report.xlsx"
    export_to_excel(storage.load_calibration_data(path), str(output))
    wb = load_workbook(output)
    try:
        values = {row[0]: row[1] for row in wb["校准信息"].iter_rows(min_row=3, values_only=True)}
        assert values["校准状态"] == "已中断（部分结果）"
        assert values["SCPI 失败次数"] == 2
        assert values["失败项目 amp"] == "读取失败"
    finally:
        wb.close()
