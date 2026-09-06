"""校准报告分组测试：同一校准项目的所有通道数据合并到同一 sheet / 同一表格。

注意：本环境沙箱会锁定 pytest 自管的临时目录，因此使用普通 os.makedirs
在工作区 data/ 下自建临时目录，并在测试结束后清理。
"""

import json
import os
import re
import shutil

import pytest
from click.testing import CliRunner
from openpyxl import load_workbook

from osccal.cli import cli
from osccal.core.export import export_to_excel

# 合成 4 通道全项目数据（与真实校准 JSON 结构一致）
CHANNELS = ["1", "2", "3", "4"]
_TMP_DIR = os.path.join("data", ".test_tmp")


@pytest.fixture
def workdir():
    """工作区内的临时目录（可写、gitignored）。"""
    os.makedirs(_TMP_DIR, exist_ok=True)
    yield _TMP_DIR
    shutil.rmtree(_TMP_DIR, ignore_errors=True)


def _make_row(ch: str, idx: int) -> dict:
    return {
        "index": idx,
        "channel": ch,
        "scale": 0.1,
        "std_value": 0.6,
        "measured": 0.6006,
        "error": 0.1,
    }


def _make_results() -> dict:
    results = {}
    for ch in CHANNELS:
        results[f"amp_ch{ch}"] = [_make_row(ch, i) for i in range(1, 3)]
        results[f"dc_gain_ch{ch}"] = [
            {"index": i, "channel": ch, "error": 0.2} for i in range(1, 3)
        ]
    return results


def _make_data() -> dict:
    return {
        "metadata": {
            "timestamp": "2026-08-21T00:00:00",
            "channel": "1,2,3,4",
            "probe": "9560",
            "oscilloscope": {"manufacturer": "TEK", "model": "MDO34", "serial": "S1"},
            "calibrator": {"manufacturer": "FLUKE", "model": "9500B"},
            "limits": {"amp": {"upper": 2.0, "lower": -2.0}},
        },
        "results": _make_results(),
    }


def test_export_groups_by_item_single_sheet(workdir):
    """同一项目的所有通道应合并到同一个 sheet，而非每通道一个 sheet。"""
    output = os.path.join(workdir, "report.xlsx")
    export_to_excel(_make_data(), output)

    wb = load_workbook(output)
    assert wb.sheetnames == ["校准信息", "幅度(ΔV)", "直流增益"], wb.sheetnames

    # 幅度 sheet：4 通道 × 2 点 = 8 行数据 + 1 行表头
    ws = wb["幅度(ΔV)"]
    assert ws.max_row == 9, f"幅度 sheet 应有 9 行，实际 {ws.max_row}"
    channels_in_sheet = {ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)}
    assert channels_in_sheet == {"1", "2", "3", "4"}, "同一 sheet 应包含全部通道数据"


def test_export_single_channel_unchanged(workdir):
    """单通道数据（无 _ch 后缀键）导出行为不变。"""
    data = _make_data()
    data["results"] = {"amp": [_make_row("1", i) for i in range(1, 3)]}
    data["metadata"]["channel"] = "1"

    output = os.path.join(workdir, "single.xlsx")
    export_to_excel(data, output)

    wb = load_workbook(output)
    assert wb.sheetnames == ["校准信息", "幅度(ΔV)"]
    assert wb["幅度(ΔV)"].max_row == 3  # 表头 + 2 行


def test_show_merges_channels_into_one_table(workdir):
    """show 命令：同一项目只渲染一个表格，包含所有通道。"""
    path = os.path.join(workdir, "data.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(_make_data(), f, ensure_ascii=False)

    runner = CliRunner()
    result = runner.invoke(cli, ["show", "--file", path])
    assert result.exit_code == 0, result.output

    # 幅度表格标题只出现一次（合并后只有一个表）
    assert result.output.count("ΔV(幅度) 校准结果") == 1
    # 合并后的表格应包含全部 4 个通道（取幅度表格区段，解析第 2 列的通道值）
    amp_section = result.output.split("ΔV(幅度) 校准结果")[1].split("直流增益 校准结果")[0]
    channels = set(re.findall(r"│\s*\d+\s*│\s*([1-4])\s*│", amp_section))
    assert channels == set(CHANNELS), f"合并表格应含全部通道，实际 {channels}"


def test_show_legacy_multi_channel_file(workdir):
    """旧版多通道数据（无 limits 字段）兼容性：不抛异常且正常渲染。"""
    data = _make_data()
    data["metadata"].pop("limits", None)  # 模拟旧数据
    path = os.path.join(workdir, "legacy.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)

    runner = CliRunner()
    result = runner.invoke(cli, ["show", "--file", path])
    assert result.exit_code == 0, result.output
    assert result.output.count("直流增益 校准结果") == 1
