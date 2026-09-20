"""导出/回看的显示精度：Excel 与 show 表格的有效位数与校准时的终端表格一致。

读数直接来自 SCPI（如 0.01232000000000001、7.96248e-08），若不整理，报告里会
出现大量无意义的位数。这里守护 `table_configs.ITEM_PRECISION` 与两个渲染入口
（`export.py`、`cli._display_calibration_data`）之间的契约。

注意：本环境沙箱会锁定 pytest 自管的临时目录，故在工作区 data/ 下自建临时目录。
"""

import json
import os
import shutil

import pytest
from click.testing import CliRunner
from openpyxl import load_workbook

from osccal.cli import cli
from osccal.core.export import export_to_excel
from osccal.core.table_configs import EXCEL_ITEM_CONFIGS, ITEM_CONFIGS, ITEM_PRECISION
from osccal.core.utils import (
    fixed_precision_decimals,
    format_display_value,
    format_with_fixed_precision,
    prepare_excel_value,
)

_TMP_DIR = os.path.join("data", ".test_tmp_precision")


@pytest.fixture
def workdir():
    os.makedirs(_TMP_DIR, exist_ok=True)
    yield _TMP_DIR
    shutil.rmtree(_TMP_DIR, ignore_errors=True)


def _export(workdir, results: dict, name: str = "precision.xlsx") -> str:
    output = os.path.join(workdir, name)
    export_to_excel({"metadata": {}, "results": results}, output)
    return output


# ── 规格表本身 ───────────────────────────────────────────────────────────


def test_precision_table_covers_every_column_of_both_configs():
    """精度规格必须与终端、Excel 两份配置的列数一一对应，改列时不能漏改精度。"""
    for item, specs in ITEM_PRECISION.items():
        assert (
            len(specs)
            == len(ITEM_CONFIGS[item]["columns"])
            == len(EXCEL_ITEM_CONFIGS[item]["headers"])
        ), f"{item} 的精度规格与列数不一致"
        for spec in specs:
            assert spec is None or (spec[0] in {"sig", "dec"} and spec[1] >= 0)


@pytest.mark.parametrize(
    "value,precision",
    [(0.01232, 3), (0.6, 3), (30.0, 3), (0.012, 3), (7.96248e-08, 4), (3.985, 2), (1234.0, 3)],
)
def test_sig_spec_reproduces_terminal_digits(value, precision):
    """("sig", n) 的 Excel 显示位数必须与终端 format_with_fixed_precision 完全一致。"""
    excel_value, number_format = prepare_excel_value(value, ("sig", precision))
    decimals = fixed_precision_decimals(value, precision)
    assert number_format == ("0" if decimals == 0 else "0." + "0" * decimals)
    assert f"{excel_value:.{decimals}f}" == format_with_fixed_precision(value, precision)


def test_non_numeric_and_spec_free_values_pass_through():
    assert prepare_excel_value("1MΩ", None) == ("1MΩ", None)
    assert prepare_excel_value("下界（未定位交叉点）", None) == ("下界（未定位交叉点）", None)
    # 文本列即使配了精度也不做数值处理
    assert prepare_excel_value("CH1", ("sig", 3)) == ("CH1", None)
    assert format_display_value("CH1", ("sig", 3)) == "CH1"


# ── Excel 导出 ──────────────────────────────────────────────────────────


def test_amp_readings_exported_with_terminal_digits(workdir):
    """幅度读数按 3 位有效数字导出，误差保留两位小数。"""
    rows = [
        {
            "index": 1,
            "channel": "1",
            "scale": 0.002,
            "std_value": 0.012,
            "measured": 0.01232000000000001,
            "error": 2.666666666666659,
        }
    ]
    ws = load_workbook(_export(workdir, {"amp": rows}))["幅度(ΔV)"]
    assert [ws.cell(2, c).value for c in range(3, 7)] == [0.002, 0.012, 0.0123, 2.67]
    assert [ws.cell(2, c).number_format for c in range(3, 7)] == [
        "0.000",
        "0.0000",
        "0.0000",
        "0.00",
    ]


def test_delta_time_readings_exported_with_terminal_digits(workdir):
    """Δt 读数按 4 位有效数字导出（终端为定点小数的位数随量级变化）。"""
    rows = [
        {
            "index": 1,
            "channel": "1",
            "scale": 1e-8,
            "std_value": 8e-8,
            "measured": 7.96248e-8,
            "error": -0.4689999999999991,
        }
    ]
    ws = load_workbook(_export(workdir, {"delta_time": rows}))["Δt(时间)"]
    assert ws.cell(2, 4).number_format == "0." + "0" * 11
    assert f"{ws.cell(2, 5).value:.11f}" == "0.00000007962"
    assert ws.cell(2, 6).value == -0.47


def test_dc_gain_and_transient_columns_use_their_own_precision(workdir):
    dc_rows = [
        {
            "index": 1,
            "channel": "1",
            "impedance": "1MΩ",
            "scale": 0.1,
            "std_value_p": 0.30000000000000004,
            "std_value_n": -0.3,
            "measured_p": 0.30123456,
            "measured_n": -0.29876543,
            "error": 0.123456,
        }
    ]
    ws = load_workbook(_export(workdir, {"dc_gain": dc_rows}, "dc.xlsx"), data_only=False)[
        "直流增益"
    ]
    assert ws.cell(2, 3).value == "1MΩ"
    assert ws.cell(2, 5).value == 0.3
    assert ws.cell(2, 7).value == 0.301
    assert ws.cell(2, 9).value == 0.12

    tr_rows = [
        {"channel": "1", "risetime_ns": 3.985, "pos_overshoot": 4.567, "edge_speed_ps": 500.0}
    ]
    ws = load_workbook(_export(workdir, {"transient": tr_rows}, "tr.xlsx"), data_only=False)[
        "上升时间及过冲"
    ]
    assert [ws.cell(2, c).value for c in range(2, 5)] == [4.0, 4.6, 500]
    assert ws.cell(2, 4).number_format == "0"  # 探头上升时间按整数 ps 显示


def test_bandwidth_value_exported_with_two_decimals(workdir):
    rows = [
        {"index": 1, "channel": "1", "scale": 0.1, "bandwidth_mhz": 99.999999, "status": "测得"}
    ]
    ws = load_workbook(_export(workdir, {"bandwidth": rows}))["频带宽度"]
    assert ws.cell(2, 4).value == 100.0
    assert ws.cell(2, 4).number_format == "0.00"
    assert ws.cell(2, 5).value == "测得"


def test_rounding_does_not_hide_out_of_limit_error(workdir):
    """误差单元格按两位小数整理，但标红仍按未舍入原值（2.004% > 2%）。"""
    rows = [
        {
            "index": 1,
            "channel": "1",
            "scale": 0.1,
            "std_value": 0.6,
            "measured": 0.612,
            "error": 2.004,
        }
    ]
    output = os.path.join(workdir, "limit.xlsx")
    export_to_excel(
        {"metadata": {"limits": {"amp": {"lower": -2.0, "upper": 2.0}}}, "results": {"amp": rows}},
        output,
    )
    ws = load_workbook(output)["幅度(ΔV)"]
    assert ws.cell(2, 6).value == 2.0
    assert ws.cell(2, 6).fill.fgColor.rgb == "00FF0000"


def test_missing_precision_config_keeps_raw_values(workdir):
    """旧配置（无 precision 字段）导出行为不变，避免历史文件无法导出。"""
    from openpyxl import Workbook

    from osccal.core.export import _create_data_sheet

    wb = Workbook()
    config = {"title": "自定义", "headers": ["序号", "值"], "error_col": None}
    _create_data_sheet(wb, "custom", config, [[1, 0.123456789]])
    ws = wb["自定义"]
    assert ws.cell(2, 2).value == 0.123456789
    assert ws.cell(2, 2).number_format == "General"


# ── show 回看 ───────────────────────────────────────────────────────────


def test_show_prints_readings_with_terminal_digits(workdir):
    data = {
        "metadata": {},
        "results": {
            "amp": [
                {
                    "index": 1,
                    "channel": "1",
                    "scale": 0.002,
                    "std_value": 0.012,
                    "measured": 0.01232000000000001,
                    "error": 2.666666666666659,
                }
            ]
        },
    }
    path = os.path.join(workdir, "show.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)

    result = CliRunner().invoke(cli, ["show", "--file", path])
    assert result.exit_code == 0, result.output
    assert "0.0123" in result.output
    assert "2.67" in result.output
    assert "0.01232000000000001" not in result.output
