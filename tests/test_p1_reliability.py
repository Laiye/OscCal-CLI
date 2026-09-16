"""校准可靠性回归：边界搜索、无效读数、退出关闭输出和 Socket 覆盖。"""

import copy
import time
from unittest.mock import Mock

import pytest
from openpyxl import Workbook

from osccal.core.comm import ScpiError, read_measurement
from osccal.measure.bandwidth import BandwidthCalibrator
from osccal.measure.base import OutputShutdownError
from osccal.measure.dc_gain import DcGainCalibrator
from osccal.measure.registry import CALIBRATORS_MAP


@pytest.fixture(autouse=True)
def no_sleep(monkeypatch):
    monkeypatch.setattr(time, "sleep", lambda _: None)


@pytest.fixture
def bandwidth_cal(mdo3_commands, fluke_calibrator, mdo34_profile, fake_osc):
    return BandwidthCalibrator(
        fake_osc,
        fake_osc,
        mdo3_commands,
        fluke_calibrator,
        copy.deepcopy(mdo34_profile),
        "1",
        "9530",
    )


@pytest.mark.parametrize("mode", ["bisect", "linear"])
def test_downward_search_without_pass_terminates(bandwidth_cal, mode):
    frequencies = []

    def measure(freq):
        frequencies.append(freq)
        assert len(frequencies) < 10, "搜索未退出"
        assert 50e3 <= freq <= 600e6
        return 0.1

    bandwidth_cal._measure_amp_at = measure
    with pytest.raises(ValueError, match="未找到"):
        getattr(bandwidth_cal, f"_scan_{mode}")(1, 100e6, 20e6, 600e6, 0.1)
    assert len(frequencies) == len(set(frequencies))


@pytest.mark.parametrize("mode", ["bisect", "linear"])
@pytest.mark.parametrize("start", [595e6, 600e6])
def test_upper_limit_is_reported_as_bound(bandwidth_cal, mode, start):
    bandwidth_cal._measure_amp_at = Mock(side_effect=AssertionError("不可越界测量"))
    freq, status = getattr(bandwidth_cal, f"_scan_{mode}")(1, start, 20e6, 600e6, 1)
    assert freq == start
    assert status.startswith("下界")


@pytest.mark.parametrize("mode", ["bisect", "linear"])
def test_failed_start_at_upper_limit_searches_down(bandwidth_cal, mode):
    bandwidth_cal._measure_amp_at = lambda f: 1 if f <= 560e6 else 0.1
    freq, status = getattr(bandwidth_cal, f"_scan_{mode}")(1, 600e6, 20e6, 600e6, 0.1)
    assert freq == 560e6
    assert status == "测得"


@pytest.mark.parametrize(
    "start,step",
    [
        (700e6, 20e6),
        (0, 20e6),
        (100e6, 0),
        (100e6, -1),
        (float("nan"), 20e6),
        (100e6, float("inf")),
    ],
)
def test_invalid_scan_never_enables_output(bandwidth_cal, start, step):
    bandwidth_cal.profile.update(bandwidth=start, bd_step=step)
    bandwidth_cal._write_calibrator = Mock()
    with pytest.raises(ValueError):
        bandwidth_cal.run()
    assert all(
        call.args != ("set_output", "ON") for call in bandwidth_cal._write_calibrator.call_args_list
    )


@pytest.mark.parametrize("value", ["nan", "inf", "-inf", "9.91e37", "-9.91e37", "bad", ""])
def test_invalid_reading_raises(mdo3_commands, fake_osc, value):
    fake_osc.default_response = f":MEASUREMENT:IMMED:VALUE {value}"
    with pytest.raises(ScpiError):
        read_measurement(fake_osc, None, mdo3_commands, "1", "AMPlitude")


def test_dc_zero_gain_rejected_and_raw_precision_used(bandwidth_cal):
    cal = DcGainCalibrator(
        **{
            "inst_osc": bandwidth_cal.inst_osc,
            "inst_calibrator": bandwidth_cal.inst_calibrator,
            "cmd_osc": bandwidth_cal.cmd_osc,
            "cmd_calibrator": bandwidth_cal.cmd_calibrator,
            "profile": {"vertical_div": 8},
            "channel": "1",
            "probe": "9530",
        }
    )
    cal._adjust_vertical_position = lambda _: None
    cal._read_meas = lambda _: 0.0
    with pytest.raises(ValueError, match="增益无效"):
        cal._measure_dc_pair(0.1, "1M", 1)
    values = iter([0.3004, -0.3004])
    cal._read_meas = lambda _: next(values)
    result = cal._measure_dc_pair(0.1, "1M", 1)
    expected = 100 * (1 - 0.6008 / 0.6) / (0.6008 / 0.6)
    assert result["error"] == pytest.approx(expected)
    assert result["row_data"][-1] == round(expected, 2)


@pytest.mark.parametrize("cal_class", list(CALIBRATORS_MAP.values()))
@pytest.mark.parametrize("failure", [RuntimeError, KeyboardInterrupt])
def test_every_calibrator_disables_output_on_failure(bandwidth_cal, cal_class, failure):
    cal = cal_class(
        bandwidth_cal.inst_osc,
        bandwidth_cal.inst_calibrator,
        bandwidth_cal.cmd_osc,
        bandwidth_cal.cmd_calibrator,
        bandwidth_cal.profile,
        "1",
        "9530",
    )
    writes = []

    def write(action, *args):
        writes.append((action, *args))
        if (action, *args) == ("set_output", "ON"):
            raise failure("模拟输出开启后中断")

    cal._write_calibrator = write
    with pytest.raises(failure):
        cal.run()
    assert ("set_output", "ON") in writes
    assert writes[-1] == ("set_output", "OFF")


def test_shutdown_failure_stops_all_projects(bandwidth_cal, monkeypatch):
    from osccal.measure.all import run_all

    def broken_write(*args):
        raise ConnectionError("通信中断")

    bandwidth_cal.inst_calibrator.write = broken_write
    with pytest.raises(OutputShutdownError):
        run_all(
            bandwidth_cal.inst_osc,
            bandwidth_cal.inst_calibrator,
            bandwidth_cal.cmd_osc,
            bandwidth_cal.cmd_calibrator,
            bandwidth_cal.profile,
            "1",
            "9530",
        )


def test_socket_override_drives_socket_measurement(mdo3_commands, monkeypatch):
    import osccal.cli as cli_mod

    class Socket:
        def __init__(self):
            self.sent = []

        def sendall(self, data):
            self.sent.append(data)

        def recv(self, size):
            return b":MEASUREMENT:IMMED:VALUE 1.5\n"

    sock = Socket()
    monkeypatch.setattr(cli_mod, "connect_socket", lambda *args: (sock, {}))
    commands = copy.deepcopy(mdo3_commands)
    inst, _ = cli_mod._connect_oscilloscope(commands, None, "localhost:5025", [])
    assert read_measurement(inst, None, commands, "1", "AMPlitude") == 1.5
    assert len(sock.sent) == 3


def test_bandwidth_bound_is_explicit_in_report():
    from osccal.core.export import _create_data_sheet
    from osccal.core.table_configs import EXCEL_ITEM_CONFIGS

    wb = Workbook()
    row = {
        "index": 1,
        "channel": "1",
        "scale": 0.1,
        "bandwidth_mhz": 600,
        "status": "下界（未定位交叉点）",
    }
    _create_data_sheet(
        wb, "bandwidth", EXCEL_ITEM_CONFIGS["bandwidth"], [row], {"bandwidth": {"min_mhz": 1000}}
    )
    ws = wb["频带宽度"]
    assert ws["E1"].value == "结果状态"
    assert ws["E2"].value.startswith("下界")
    assert ws["D2"].fill.fill_type is None
