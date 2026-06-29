import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.utils import get_column_letter
from rich.console import Console
from osccal.core.table_configs import EXCEL_ITEM_CONFIGS as item_configs

console = Console()

def export_to_excel(data: dict, output_path: str) -> str:
    wb = Workbook()

    metadata = data.get("metadata", {})
    results = data.get("results", {})

    _create_info_sheet(wb, metadata)

    for item_name, rows in results.items():
        if not rows:
            continue
        base_name = item_name.split("_ch")[0]
        config = item_configs.get(base_name)
        if config:
            sheet_title = config["title"]
            if "_ch" in item_name:
                ch = item_name.split("_ch")[1]
                sheet_title = f"{config['title']} CH{ch}"
            _create_data_sheet(wb, item_name, config, rows, sheet_title)

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    wb.save(output_path)
    console.print(f"[green]✓[/green] 报告已导出: [cyan]{output_path}[/cyan]")
    return output_path


def _create_info_sheet(wb, metadata):
    ws = wb.active
    ws.title = "校准信息"

    header_font = Font(bold=True, size=14)
    label_font = Font(bold=True, size=11)
    normal_font = Font(size=11)
    border = Border(
        bottom=Side(style="thin")
    )

    ws["A1"] = "示波器校准报告"
    ws["A1"].font = Font(bold=True, size=18)
    ws.merge_cells("A1:D1")

    row = 3
    info_items = [
        ("校准时间", metadata.get("timestamp", "")),
        ("通道", f"CH{metadata.get('channel', '')}"),
    ]

    osc = metadata.get("oscilloscope", {})
    info_items.extend([
        ("示波器厂家", osc.get("manufacturer", "")),
        ("示波器型号", osc.get("model", "")),
        ("示波器序列号", osc.get("serial", "")),
        ("示波器固件版本", osc.get("firmware", "")),
    ])

    cal = metadata.get("calibrator", {})
    info_items.extend([
        ("校准仪厂家", cal.get("manufacturer", "")),
        ("校准仪型号", cal.get("model", "")),
    ])

    for label, value in info_items:
        ws.cell(row=row, column=1, value=label).font = label_font
        ws.cell(row=row, column=2, value=value).font = normal_font
        ws.cell(row=row, column=1).border = border
        ws.cell(row=row, column=2).border = border
        row += 1

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 40


def _create_data_sheet(wb, item_name, config, rows, sheet_title=None):
    title = sheet_title or config["title"]
    ws = wb.create_sheet(title=title)

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx, header in enumerate(config["headers"], 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    red_font = Font(color="FFFFFF", bold=True)

    for row_idx, row_data in enumerate(rows, 2):
        if isinstance(row_data, dict):
            values = list(row_data.values())
        else:
            values = row_data

        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = center_align
            cell.border = thin_border

            error_col = config.get("error_col")
            if error_col is not None and col_idx == error_col + 1:
                try:
                    if abs(float(val)) > 2.0:
                        cell.fill = red_fill
                        cell.font = red_font
                except (ValueError, TypeError):
                    pass

    for col_idx in range(1, len(config["headers"]) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    chart_config = config.get("chart_config")
    if chart_config and len(rows) > 1:
        _add_chart(ws, chart_config, len(rows))


def _add_chart(ws, chart_config, data_rows):
    chart = ScatterChart()
    chart.title = chart_config["title"]
    chart.x_axis.title = chart_config["x_label"]
    chart.y_axis.title = chart_config["y_label"]
    chart.style = 13

    x_col = chart_config["x_col"] + 1
    y_col = chart_config["y_col"] + 1

    x_values = Reference(ws, min_col=x_col, min_row=2, max_row=data_rows + 1)
    y_values = Reference(ws, min_col=y_col, min_row=2, max_row=data_rows + 1)

    series = Series(y_values, x_values, title=chart_config["title"])
    chart.series.append(series)

    chart.width = 20
    chart.height = 12

    ws.add_chart(chart, f"A{data_rows + 4}")
